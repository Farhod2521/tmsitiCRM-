"""Bugungi (yoki tanlangan kun) xodimlar davomati ro'yxati — Davomat sahifasidagi
"Bugun" yorlig'i uchun: kelish vaqti, holat, kechikish, ariza/holat (status) va
xodim rasmining kichik nusxasi."""
import base64
import io
from datetime import date as date_cls, datetime
from typing import List, Optional

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from PIL import Image
from pydantic import BaseModel
from sqlalchemy.orm import Session

from .. import models
from ..database import get_db
from ..deps import get_current_employee
from ..status_periods import status_code_map
from .attendance import TZ_UZ, APPROVED_NOTE_STATUSES, late_minutes_for
from .holidays import holiday_map

router = APIRouter(prefix="/attendance/daily", tags=["Kunlik davomat"])

R = models.RoleEnum
_VIEW_ROLES = {R.kadr, R.superadmin, R.direktor, R.zamdirektor}
_EXCLUDED_ROLES = {R.superadmin, R.direktor, R.zamdirektor}
_EXCLUDED_STATUSES = {models.EmployeeStatusEnum.shafyor_farrosh, models.EmployeeStatusEnum.dekret}

STATUS_CODE_LABEL = {"MT": "Mehnat ta'tili", "B": "Bolnichniy", "K": "Xizmat safari", "O'": "O'quv ta'tili"}
EMP_STATUS_LABEL = {
    "faol": "Faol", "otpuska": "Mehnat ta'tilida", "dekret": "Dekretda", "shafyor_farrosh": "Texnik xodim",
    "xizmat_safarida": "Xizmat safarida", "oquv_tatilida": "O'quv ta'tilida",
    "mehnatga_layoqatsiz": "Bolnichniy", "online": "Online",
}
NOTE_LABEL = {"kechikish": "Kechikaman", "kelmaslik": "Kelmayman", "obyektda": "Obyektda", "ruxsat": "Ruxsat"}

# Tartib: kechikkanlar tepada (ko'p kechikkan birinchi), keyin kelganlar (erta kelgan birinchi)...
_HOLAT_ORDER = {"kechikkan": 0, "kelgan": 1, "sababli": 2, "tatilda": 3, "kelmagan": 4}


class DailyNote(BaseModel):
    type: str
    label: str
    status: str
    text: Optional[str] = None


class DailyRow(BaseModel):
    employee_id: int
    full_name: str
    position: Optional[str] = None
    department: Optional[str] = None
    avatar: Optional[str] = None          # kichik rasm (data:image/jpeg;base64,...)
    check_in: Optional[str] = None        # "08:12"
    late_min: int = 0
    holat: str                            # kelgan | kechikkan | kelmagan | sababli | tatilda
    holat_label: str
    status: str                           # xodim holati: faol, otpuska, ... (yoki ariza turi)
    status_label: str
    note: Optional[DailyNote] = None
    distance_m: Optional[float] = None
    turniket_check_in: Optional[str] = None   # turniketdagi kirish vaqti (bo'lsa)
    corrected: bool = False                    # kelish vaqti superadmin tomonidan tuzatilgan


class DailyOut(BaseModel):
    date: str
    day_off: Optional[str] = None         # "Dam olish kuni" / bayram nomi
    rows: List[DailyRow]


# ── Rasm kichik nusxasi (xotirada keshlanadi) ─────────────────────────────────
_THUMB_CACHE: dict[tuple[int, int, str], Optional[str]] = {}


def _thumb(emp: models.Employee) -> Optional[str]:
    src = emp.photo_base64
    if not src:
        return None
    key = (emp.id, len(src), src[-32:])
    if key in _THUMB_CACHE:
        return _THUMB_CACHE[key]
    try:
        raw = base64.b64decode(src.split(",", 1)[1] if src.startswith("data:") else src)
        img = Image.open(io.BytesIO(raw)).convert("RGB")
        w, h = img.size
        side = min(w, h)
        img = img.crop(((w - side) // 2, (h - side) // 4, (w - side) // 2 + side, (h - side) // 4 + side))
        img.thumbnail((72, 72))
        buf = io.BytesIO()
        img.save(buf, "JPEG", quality=75)
        out = "data:image/jpeg;base64," + base64.b64encode(buf.getvalue()).decode()
    except Exception:
        out = None
    if len(_THUMB_CACHE) > 2000:
        _THUMB_CACHE.clear()
    _THUMB_CACHE[key] = out
    return out


# ── Hisob ─────────────────────────────────────────────────────────────────────

def _build(db: Session, day: str, with_avatars: bool = True) -> DailyOut:
    try:
        d = datetime.strptime(day, "%Y-%m-%d").date()
    except ValueError:
        raise HTTPException(status_code=400, detail="Sana formati noto'g'ri (YYYY-MM-DD)")

    employees = (
        db.query(models.Employee)
        .filter(models.Employee.role.notin_(list(_EXCLUDED_ROLES)),
                models.Employee.status.notin_(list(_EXCLUDED_STATUSES)))
        .all()
    )
    atts = {a.employee_id: a for a in db.query(models.Attendance).filter(models.Attendance.date == day).all()}
    codes = status_code_map(db, employees, day, day)
    notes: dict[int, models.AttendanceNote] = {}
    for n in (db.query(models.AttendanceNote)
              .filter(models.AttendanceNote.date_from <= day, models.AttendanceNote.date_to >= day)
              .order_by(models.AttendanceNote.created_at).all()):
        notes[n.employee_id] = n        # eng oxirgisi
    turniket = {t.employee_id: t.check_in for t in db.query(models.TurniketAttendance)
                .filter(models.TurniketAttendance.date == day).all() if t.check_in}
    corrected_ids = {c.employee_id for c in db.query(models.AttendanceCorrection.employee_id)
                     .filter(models.AttendanceCorrection.date == day).all()}
    holidays = holiday_map(db, day, day)
    day_off = holidays.get(day) or ("Dam olish kuni" if d.weekday() >= 5 else None)

    rows: list[DailyRow] = []
    for e in employees:
        att = atts.get(e.id)
        note = notes.get(e.id)
        approved = note is not None and note.review_status in APPROVED_NOTE_STATUSES
        code = codes.get((e.id, day))
        late = 0
        check_in = None
        if att is not None:
            ci = att.check_in.astimezone(TZ_UZ) if att.check_in.tzinfo is not None else att.check_in
            check_in = ci.strftime("%H:%M")
            late = late_minutes_for(ci)
            if late > 0 and approved:
                holat, holat_label = "kelgan", "Kelgan"          # sababli kechikish
            elif late > 0:
                holat, holat_label = "kechikkan", "Kechikkan"
            else:
                holat, holat_label = "kelgan", "Kelgan"
        elif code:
            holat, holat_label = "tatilda", STATUS_CODE_LABEL.get(code, code)
        elif approved:
            holat, holat_label = "sababli", "Sababli"
        else:
            holat, holat_label = "kelmagan", "Kelmagan"

        # Status ustuni: kunni qamragan ariza bo'lsa — ariza turi, aks holda xodim holati
        if note is not None:
            status, status_label = f"ariza_{note.note_type}", NOTE_LABEL.get(note.note_type, note.note_type)
        elif e.status == models.EmployeeStatusEnum.faol and not e.is_active:
            status, status_label = "faol_emas", "Faol emas"
        else:
            st = e.status.value if hasattr(e.status, "value") else str(e.status)
            status, status_label = st, EMP_STATUS_LABEL.get(st, st)

        rows.append(DailyRow(
            employee_id=e.id, full_name=e.full_name, position=e.position,
            department=e.department.name if e.department else None,
            avatar=_thumb(e) if with_avatars else None,
            check_in=check_in, late_min=late if holat == "kechikkan" else 0,
            holat=holat, holat_label=holat_label, status=status, status_label=status_label,
            note=DailyNote(type=note.note_type, label=NOTE_LABEL.get(note.note_type, note.note_type),
                           status=note.review_status, text=note.text) if note else None,
            distance_m=att.distance_m if att else None,
            turniket_check_in=turniket.get(e.id),
            corrected=e.id in corrected_ids,
        ))

    rows.sort(key=lambda r: (
        _HOLAT_ORDER.get(r.holat, 9),
        -r.late_min if r.holat == "kechikkan" else 0,
        r.check_in or "99:99",
        r.full_name,
    ))
    return DailyOut(date=day, day_off=day_off, rows=rows)


def _require(current: models.Employee) -> None:
    if current.role not in _VIEW_ROLES:
        raise HTTPException(status_code=403, detail="Ruxsat yo'q")


@router.get("", response_model=DailyOut)
def daily_list(
    date: Optional[str] = None,
    db: Session = Depends(get_db),
    current: models.Employee = Depends(get_current_employee),
):
    _require(current)
    return _build(db, date or datetime.now(TZ_UZ).strftime("%Y-%m-%d"))


# ── Kelish vaqtini tuzatish (faqat superadmin) ────────────────────────────────

class CorrectionIn(BaseModel):
    employee_id: int
    date: str            # "2026-09-15"
    time: str            # "08:47"
    reason: Optional[str] = None
    source: str = "qolda"   # "turniket" | "qolda"


class CorrectionLog(BaseModel):
    old_check_in: Optional[str] = None
    new_check_in: str
    source: str
    reason: Optional[str] = None
    corrected_by: Optional[str] = None
    created_at: datetime


class CorrectionInfo(BaseModel):
    employee_id: int
    full_name: str
    date: str
    current: Optional[str] = None      # tizimdagi kelish vaqti
    turniket: Optional[str] = None     # turniketdagi kirish vaqti
    history: List[CorrectionLog]


def _require_superadmin(current: models.Employee) -> None:
    if current.role != R.superadmin:
        raise HTTPException(status_code=403, detail="Kelish vaqtini faqat superadmin tuzata oladi")


def _hhmm(dt: datetime) -> str:
    return (dt.astimezone(TZ_UZ) if dt.tzinfo is not None else dt).strftime("%H:%M")


def _correction_info(db: Session, emp: models.Employee, day: str) -> CorrectionInfo:
    att = db.query(models.Attendance).filter(models.Attendance.employee_id == emp.id, models.Attendance.date == day).first()
    tur = db.query(models.TurniketAttendance).filter(models.TurniketAttendance.employee_id == emp.id,
                                                     models.TurniketAttendance.date == day).first()
    logs = (db.query(models.AttendanceCorrection)
            .filter(models.AttendanceCorrection.employee_id == emp.id, models.AttendanceCorrection.date == day)
            .order_by(models.AttendanceCorrection.created_at.desc()).all())
    return CorrectionInfo(
        employee_id=emp.id, full_name=emp.full_name, date=day,
        current=_hhmm(att.check_in) if att else None,
        turniket=tur.check_in if tur else None,
        history=[CorrectionLog(old_check_in=l.old_check_in, new_check_in=l.new_check_in, source=l.source,
                               reason=l.reason, corrected_by=l.corrector.full_name if l.corrector else None,
                               created_at=l.created_at) for l in logs],
    )


@router.get("/correction", response_model=CorrectionInfo)
def get_correction(
    employee_id: int,
    date: str,
    db: Session = Depends(get_db),
    current: models.Employee = Depends(get_current_employee),
):
    _require_superadmin(current)
    emp = db.get(models.Employee, employee_id)
    if not emp:
        raise HTTPException(status_code=404, detail="Xodim topilmadi")
    return _correction_info(db, emp, date)


@router.put("/correction", response_model=CorrectionInfo)
def set_correction(
    data: CorrectionIn,
    db: Session = Depends(get_db),
    current: models.Employee = Depends(get_current_employee),
):
    """Kelish vaqtini o'rnatish/tuzatish. Yozuv bo'lmasa (xodim "Ishga keldim"ni
    bosmagan) — yaratiladi. Har bir o'zgarish jurnalga yoziladi."""
    _require_superadmin(current)
    emp = db.get(models.Employee, data.employee_id)
    if not emp:
        raise HTTPException(status_code=404, detail="Xodim topilmadi")
    try:
        d = datetime.strptime(data.date, "%Y-%m-%d").date()
        t = datetime.strptime(data.time.strip(), "%H:%M").time()
    except ValueError:
        raise HTTPException(status_code=400, detail="Sana (YYYY-MM-DD) yoki vaqt (HH:MM) noto'g'ri")
    if d > datetime.now(TZ_UZ).date():
        raise HTTPException(status_code=400, detail="Kelajakdagi sana uchun vaqt qo'yib bo'lmaydi")
    if data.source not in ("turniket", "qolda"):
        raise HTTPException(status_code=400, detail="Noto'g'ri manba")

    new_dt = datetime.combine(d, t)   # mahalliy (UTC+5) vaqt, naive — check_in bilan bir xil
    att = db.query(models.Attendance).filter(models.Attendance.employee_id == emp.id,
                                             models.Attendance.date == data.date).first()
    old = _hhmm(att.check_in) if att else None
    if old == t.strftime("%H:%M"):
        return _correction_info(db, emp, data.date)
    if att:
        att.check_in = new_dt
    else:
        db.add(models.Attendance(employee_id=emp.id, date=data.date, check_in=new_dt,
                                 latitude=0.0, longitude=0.0, distance_m=None))
    db.add(models.AttendanceCorrection(
        employee_id=emp.id, date=data.date, old_check_in=old, new_check_in=t.strftime("%H:%M"),
        source=data.source, reason=(data.reason or "").strip() or None, corrected_by=current.id,
    ))
    db.commit()
    return _correction_info(db, emp, data.date)


@router.get("/xlsx")
def daily_xlsx(
    date: Optional[str] = None,
    db: Session = Depends(get_db),
    current: models.Employee = Depends(get_current_employee),
):
    _require(current)
    day = date or datetime.now(TZ_UZ).strftime("%Y-%m-%d")
    data = _build(db, day, with_avatars=False)

    wb = Workbook()
    ws = wb.active
    ws.title = day
    y, m, dd = day.split("-")
    ws.merge_cells("A1:G1")
    ws["A1"] = f"Xodimlar davomati — {dd}.{m}.{y}" + (f" ({data.day_off})" if data.day_off else "")
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([])
    ws.append(["№", "Xodim", "Bo'lim", "Kelish vaqti", "Holat", "Kechikish", "Status"])
    thin = Side(style="thin", color="FFD9E3F0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for c in ws[3]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="FFF4F9FD")
        c.border = border
        c.alignment = Alignment(horizontal="center", vertical="center")
    fills = {"kechikkan": "FFFFF3CD", "kelmagan": "FFFDE2E2", "kelgan": "FFE3F7EC", "sababli": "FFEFF1F5", "tatilda": "FFE3EEFF"}
    for i, r in enumerate(data.rows, 1):
        ws.append([i, r.full_name, r.department or "", r.check_in or "—", r.holat_label,
                   f"{r.late_min} daq" if r.late_min else "—", r.status_label])
        for c in ws[ws.max_row]:
            c.border = border
            c.alignment = Alignment(horizontal="left" if c.column in (2, 3) else "center", vertical="center")
        ws.cell(row=ws.max_row, column=5).fill = PatternFill("solid", fgColor=fills.get(r.holat, "FFFFFFFF"))
    for col, w in zip("ABCDEFG", (5, 28, 34, 13, 16, 12, 16)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A4"

    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return StreamingResponse(
        buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="davomat_{day}.xlsx"'},
    )
