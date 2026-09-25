"""Ish staji — kadr xlsx'dan yuklaydi (masalan, "avgust oy uchun"), keyingi oylar
uchun staj har oyga +1 oy qo'shib hisoblanadi va shunga qarab ustama foizi
avtomatik belgilanadi. Oyning barcha ish kunlari mehnat ta'tili (MT) bo'lsa —
o'sha oy uchun foiz 0."""
import io
import re
from calendar import monthrange
from datetime import date, datetime
from typing import List, Optional

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Cm
from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, Side
from pydantic import BaseModel
from sqlalchemy.orm import Session

from .. import models
from ..database import get_db
from ..deps import get_current_employee
from ..name_match import match_employee
from ..status_periods import status_code_map
from .holidays import month_holidays

router = APIRouter(prefix="/ish-staji", tags=["Ish staji"])

_ROLES = {models.RoleEnum.kadr, models.RoleEnum.superadmin}

MONTHS_UZ = ["yanvar", "fevral", "mart", "aprel", "may", "iyun",
             "iyul", "avgust", "sentabr", "oktabr", "noyabr", "dekabr"]

# (shu oydan boshlab, foiz) — "1 yildan 3 yilgacha 25%", ..., "20 yildan ortiq 150%"
_BRACKETS = [(20 * 12, 150), (15 * 12, 125), (10 * 12, 100), (5 * 12, 75), (3 * 12, 50), (1 * 12, 25)]


def percent_for(total_months: int) -> int:
    for threshold, pct in _BRACKETS:
        if total_months >= threshold:
            return pct
    return 0


def _require(current: models.Employee) -> None:
    if current.role not in _ROLES:
        raise HTTPException(status_code=403, detail="Ruxsat yo'q")


# ── Sxemalar ──────────────────────────────────────────────────────────────────

class ExperienceRow(BaseModel):
    id: int
    order_num: Optional[int] = None
    employee_id: Optional[int] = None
    full_name: str
    position: Optional[str] = None
    years: int
    months: int
    percent: int              # shu oy uchun yakuniy foiz (MT bo'lsa 0)
    base_percent: int         # stajga ko'ra foiz
    full_month_mt: bool       # butun oy mehnat ta'tilida


class ExperienceOut(BaseModel):
    year: int
    month: int
    base_year: Optional[int] = None
    base_month: Optional[int] = None
    rows: List[ExperienceRow]


class ExperienceEditIn(BaseModel):
    year: int                 # qaysi oy holatida kiritilmoqda
    month: int
    years: int
    months: int
    position: Optional[str] = None


class ImportOut(BaseModel):
    imported: int
    matched: int
    unmatched: List[str]
    base_year: int
    base_month: int


# ── Hisob ─────────────────────────────────────────────────────────────────────

def _month_index(y: int, m: int) -> int:
    return y * 12 + (m - 1)


def _full_month_mt(db: Session, rows: list, year: int, month: int) -> set[int]:
    """Oyning BARCHA ish kunlari (dush-juma, bayramlarsiz) MT bo'lgan xodimlar id'lari —
    holatlar tarixi va kadr qo'lda tuzatgan kunlar (tabel) bo'yicha."""
    emps = [r.employee for r in rows if r.employee is not None]
    if not emps:
        return set()
    days = monthrange(year, month)[1]
    start, end = f"{year:04d}-{month:02d}-01", f"{year:04d}-{month:02d}-{days:02d}"
    codes = status_code_map(db, emps, start, end)
    overrides = {
        (o.employee_id, o.date): o.code
        for o in db.query(models.TabelOverride).filter(
            models.TabelOverride.employee_id.in_([e.id for e in emps]),
            models.TabelOverride.date >= start, models.TabelOverride.date <= end,
        ).all()
    }
    holidays = month_holidays(db, year, month)
    workdays = [date(year, month, d).isoformat() for d in range(1, days + 1)
                if date(year, month, d).weekday() < 5 and d not in holidays]
    result = set()
    for e in emps:
        if workdays and all(overrides.get((e.id, d), codes.get((e.id, d))) == "MT" for d in workdays):
            result.add(e.id)
    return result


def _build(db: Session, year: int, month: int) -> ExperienceOut:
    rows = (
        db.query(models.WorkExperience)
        .order_by(models.WorkExperience.order_num.nullslast(), models.WorkExperience.id)
        .all()
    )
    if not rows:
        return ExperienceOut(year=year, month=month, rows=[])
    mt = _full_month_mt(db, rows, year, month)
    out = []
    for r in rows:
        total = max(0, r.base_total_months + _month_index(year, month) - _month_index(r.base_year, r.base_month))
        base_pct = percent_for(total)
        full_mt = r.employee_id in mt
        out.append(ExperienceRow(
            id=r.id, order_num=r.order_num, employee_id=r.employee_id,
            full_name=r.full_name, position=r.position,
            years=total // 12, months=total % 12,
            percent=0 if full_mt else base_pct, base_percent=base_pct, full_month_mt=full_mt,
        ))
    return ExperienceOut(year=year, month=month, base_year=rows[0].base_year, base_month=rows[0].base_month, rows=out)


# ── Import ────────────────────────────────────────────────────────────────────

def _num(v) -> int:
    if v is None or v == "":
        return 0
    try:
        return int(float(str(v).replace(",", ".").strip()))
    except ValueError:
        return 0


def _norm(v) -> str:
    return re.sub(r"\s+", " ", str(v or "")).strip().lower()


def _parse_xlsx(content: bytes) -> tuple[list[dict], Optional[int]]:
    """Jadvalni topadi: "FISh" sarlavhali qator, uning ostidagi "yil"/"oy" qatori.
    Qaytaradi: (qatorlar, sarlavhadagi oy raqami yoki None)."""
    ws = load_workbook(io.BytesIO(content), data_only=True).active
    grid = [list(r) for r in ws.iter_rows(values_only=True)]

    title_month = None
    for r in grid[:15]:
        for v in r:
            m = re.search(r"([a-zA-Z']+)\s+oy(i)?\s+uchun", str(v or ""), re.I)
            if m and m.group(1).lower() in MONTHS_UZ:
                title_month = MONTHS_UZ.index(m.group(1).lower()) + 1

    hdr = next((i for i, r in enumerate(grid) if any(_norm(v) in ("fish", "f.i.sh", "f.i.sh.", "fio") for v in r)), None)
    if hdr is None:
        raise HTTPException(status_code=400, detail="Jadvalda \"FISh\" ustuni topilmadi")
    head = [_norm(v) for v in grid[hdr]]
    sub = [_norm(v) for v in grid[hdr + 1]] if hdr + 1 < len(grid) else []

    def col(names, row):
        return next((i for i, v in enumerate(row) if v in names), None)

    c_fish = col(("fish", "f.i.sh", "f.i.sh.", "fio"), head)
    c_tr = col(("tr", "t/r", "№", "n"), head)
    c_pos = next((i for i, v in enumerate(head) if v.startswith("lavozim")), None)
    c_yil = col(("yil",), sub) if col(("yil",), sub) is not None else col(("yil",), head)
    c_oy = col(("oy",), sub) if col(("oy",), sub) is not None else col(("oy",), head)
    if c_yil is None or c_oy is None:
        raise HTTPException(status_code=400, detail="\"yil\" va \"oy\" ustunlari topilmadi")

    start = hdr + (2 if col(("yil",), sub) is not None else 1)
    rows = []
    for r in grid[start:]:
        name = str(r[c_fish] or "").strip() if c_fish < len(r) else ""
        if not name:
            if rows:
                break
            continue
        rows.append({
            "order_num": _num(r[c_tr]) if c_tr is not None and c_tr < len(r) and r[c_tr] not in (None, "") else None,
            "full_name": name,
            "position": str(r[c_pos]).strip() if c_pos is not None and c_pos < len(r) and r[c_pos] else None,
            "years": _num(r[c_yil]) if c_yil < len(r) else 0,
            "months": _num(r[c_oy]) if c_oy < len(r) else 0,
        })
    if not rows:
        raise HTTPException(status_code=400, detail="Jadvalda xodimlar topilmadi")
    return rows, title_month


# ── Endpointlar ───────────────────────────────────────────────────────────────

@router.get("", response_model=ExperienceOut)
def list_experience(
    year: int,
    month: int,
    db: Session = Depends(get_db),
    current: models.Employee = Depends(get_current_employee),
):
    _require(current)
    return _build(db, year, month)


@router.post("/import", response_model=ImportOut)
async def import_experience(
    file: UploadFile = File(...),
    year: Optional[int] = Form(None),
    month: Optional[int] = Form(None),
    db: Session = Depends(get_db),
    current: models.Employee = Depends(get_current_employee),
):
    """Xlsx'ni yuklaydi va BARCHA oldingi yozuvlarni almashtiradi. year/month —
    jadval qaysi oy holatida ("avgust oy uchun"); berilmasa sarlavhadan olinadi."""
    _require(current)
    if not (file.filename or "").lower().endswith(".xlsx"):
        raise HTTPException(status_code=400, detail="Faqat .xlsx fayl qabul qilinadi")
    parsed, title_month = _parse_xlsx(await file.read())
    month = month or title_month
    if not month:
        raise HTTPException(status_code=400, detail="Jadval qaysi oy uchun ekanini tanlang")
    year = year or date.today().year

    employees = db.query(models.Employee).all()
    db.query(models.WorkExperience).delete()
    matched, unmatched = 0, []
    for p in parsed:
        emp, _conf = match_employee(p["full_name"], employees)
        if emp:
            matched += 1
        else:
            unmatched.append(p["full_name"])
        db.add(models.WorkExperience(
            order_num=p["order_num"], employee_id=emp.id if emp else None,
            full_name=p["full_name"], position=p["position"],
            base_year=year, base_month=month,
            base_total_months=p["years"] * 12 + p["months"], updated_by=current.id,
        ))
    db.commit()
    return ImportOut(imported=len(parsed), matched=matched, unmatched=unmatched, base_year=year, base_month=month)


@router.patch("/{row_id}", response_model=ExperienceRow)
def edit_experience(
    row_id: int,
    data: ExperienceEditIn,
    db: Session = Depends(get_db),
    current: models.Employee = Depends(get_current_employee),
):
    """Stajni tanlangan oy holatida tuzatish — asos oyga qayta hisoblab saqlanadi."""
    _require(current)
    r = db.query(models.WorkExperience).filter(models.WorkExperience.id == row_id).first()
    if not r:
        raise HTTPException(status_code=404, detail="Topilmadi")
    if data.years < 0 or not 0 <= data.months <= 11:
        raise HTTPException(status_code=400, detail="Yil 0 dan katta, oy 0–11 oralig'ida bo'lishi kerak")
    total = data.years * 12 + data.months
    r.base_total_months = total - (_month_index(data.year, data.month) - _month_index(r.base_year, r.base_month))
    if data.position is not None:
        r.position = data.position.strip() or None
    r.updated_by = current.id
    db.commit()
    return next(x for x in _build(db, data.year, data.month).rows if x.id == row_id)


# ── Yuklab olish ──────────────────────────────────────────────────────────────

_HEADER_TO = "Buxgalteriya va moliya bo'limiga"
_HEADER_FROM = "Inson resurslarini rivojlantirish bo'limidan"
_ORG = "Texnik me’yorlash va standartlashtirish ilmiy-tadqiqot instituti"


def _title(month: int) -> str:
    return f"{_ORG} xodimlarining {MONTHS_UZ[month - 1]} oy uchun ish stajlari to'grisida"


@router.get("/xlsx")
def download_xlsx(
    year: int,
    month: int,
    db: Session = Depends(get_db),
    current: models.Employee = Depends(get_current_employee),
):
    _require(current)
    data = _build(db, year, month)
    wb = Workbook()
    ws = wb.active
    ws.title = f"{MONTHS_UZ[month - 1]} {year}"[:31]
    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws.merge_cells("D1:F1"); ws["D1"] = _HEADER_TO; ws["D1"].font = Font(bold=True, size=12)
    ws.merge_cells("D2:F2"); ws["D2"] = _HEADER_FROM; ws["D2"].font = Font(size=12)
    for c in ("D1", "D2"):
        ws[c].alignment = center
    ws.merge_cells("A4:F4"); ws["A4"] = _title(month)
    ws["A4"].font = Font(bold=True, size=12); ws["A4"].alignment = center
    ws.row_dimensions[4].height = 36
    ws.merge_cells("A5:F5"); ws["A5"] = "MA'LUMOT"; ws["A5"].font = Font(bold=True, size=12); ws["A5"].alignment = center

    # Sarlavha: Tr | FISh | Lavozimi | Ish staji (yil, oy) | Foiz
    for rng, val in (("A7:A8", "Tr"), ("B7:B8", "FISh"), ("C7:C8", "Lavozimi"), ("D7:E7", "Ish staji"), ("F7:F8", "Foiz")):
        ws.merge_cells(rng)
        ws[rng.split(":")[0]] = val
    ws["D8"], ws["E8"] = "yil", "oy"
    for row in ws.iter_rows(min_row=7, max_row=8, min_col=1, max_col=6):
        for c in row:
            c.font = Font(bold=True); c.alignment = center; c.border = border

    for i, r in enumerate(data.rows, start=1):
        ws.append([r.order_num or i, r.full_name, r.position or "", r.years, r.months, r.percent])
        for c in ws[ws.max_row]:
            c.border = border
            c.alignment = Alignment(horizontal="left" if c.column == 2 else "center", vertical="center", wrap_text=True)

    for col, w in zip("ABCDEF", (6, 26, 30, 8, 8, 10)):
        ws.column_dimensions[col].width = w

    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return StreamingResponse(
        buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="ish_staji_{year}_{month:02d}.xlsx"'},
    )


@router.get("/docx")
def download_docx(
    year: int,
    month: int,
    db: Session = Depends(get_db),
    current: models.Employee = Depends(get_current_employee),
):
    _require(current)
    data = _build(db, year, month)
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.left_margin = sec.right_margin = Cm(2)
    sec.top_margin = sec.bottom_margin = Cm(1.5)
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    def para(text, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER, size=12, indent_cm=None, space_after=0):
        p = doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_after = Pt(space_after)
        if indent_cm:
            p.paragraph_format.left_indent = Cm(indent_cm)
        run = p.add_run(text)
        run.bold = bold
        run.font.size = Pt(size)
        return p

    para(_HEADER_TO, bold=True, indent_cm=9.5)
    para(_HEADER_FROM, indent_cm=9.5, space_after=18)
    para(_title(month), bold=True, size=13)
    para("MA'LUMOT", bold=True, size=13, space_after=10)

    t = doc.add_table(rows=2, cols=6)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    h1, h2 = t.rows[0].cells, t.rows[1].cells
    for i, txt in ((0, "Tr"), (1, "FISh"), (2, "Lavozimi"), (5, "Foiz")):
        h1[i].merge(h2[i]).text = txt
    h1[3].merge(h1[4]).text = "Ish staji"
    h2[3].text, h2[4].text = "yil", "oy"
    for row in t.rows[:2]:
        for c in row.cells:
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in c.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.bold = True

    for i, r in enumerate(data.rows, start=1):
        cells = t.add_row().cells
        vals = [str(r.order_num or i), r.full_name, r.position or "", str(r.years), str(r.months), str(r.percent)]
        for j, v in enumerate(vals):
            cells[j].text = v
            cells[j].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 1 else WD_ALIGN_PARAGRAPH.CENTER

    for row in t.rows:
        for j, w in enumerate((1.1, 4.6, 5.4, 1.5, 1.5, 1.9)):
            row.cells[j].width = Cm(w)

    buf = io.BytesIO(); doc.save(buf); buf.seek(0)
    return StreamingResponse(
        buf, media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="ish_staji_{year}_{month:02d}.docx"'},
    )
