"""Avans (yarim oylik) va bo'lim tabeli — bo'limma-bo'lim davomat jadvali.

Kunlik kod avtomatik olinadi (bitta manba — kadr ham, bo'lim boshlig'i ham bir xil ko'radi):
  dam olish kuni → X, bayram (kadr kalendari) → Б, kadr qo'lda tuzatgan kun → o'sha kod,
  mehnat ta'tili / o'quv ta'tili / bolnichniy / safar / BS (holatlar tarixidan) → М/Т, У/Т, К, Х/С, Б/С,
  dekret → Д, ishga kelgan (yoki hali kelmagan kelajak ish kuni) → 8 (yarim stavka — 4),
  o'tgan ish kuni, kelmagan va sababi belgilanmagan → bo'sh.
Bo'lim boshlig'i qo'lda qo'ygan kod (8/4 dan boshqa) — faqat qulflanmagan ish kunlariga qo'llanadi."""
import io
from calendar import monthrange
from datetime import date, datetime

from sqlalchemy.orm import Session

from . import models
from .routers.attendance import TZ_UZ
from .routers.holidays import month_holidays
from .status_periods import status_code_map

R, S = models.RoleEnum, models.EmployeeStatusEnum

# Avtomatik (ichki) kod → jadvaldagi kod
DISPLAY = {"MT": "М/Т", "O'": "У/Т", "B": "К", "K": "Х/С", "BS": "Б/С", "Д": "Д", "BY": "Б", "X": "X", "8": "8", "": ""}
LEGEND = [
    ("4", "Yarim stavka (4 soat)"), ("8", "Ish kuni (8 soat)"), ("X", "Dam olish"), ("Б", "Bayram"),
    ("К", "Kasallik varaqasi"), ("У/Т", "O'quv ta'tili"), ("М/Т", "Mehnat ta'tili"),
    ("Х/С", "Xizmat safari"), ("Б/С", "Ish haqisiz ta'til"), ("Д", "Dekret"),
]
WORK_CODES = {"8", "4", "Х/С"}                 # I.K. — ishlangan kunlar
ABSENT_CODES = {"М/Т", "У/Т", "К", "Б/С", "Д"}  # I.Km. — ishlanmagan (sababli) kunlar
_EXCLUDED_ROLES = {R.superadmin, R.direktor, R.zamdirektor}
_HEAD_ROLES = (R.bolim_boshligi, R.boshqarma_boshligi)


def _work_code(emp) -> str:
    return "4" if (emp.work_rate or 1) <= 0.5 else "8"


def build(db: Session, year: int, month: int, day_from: int = 1, day_to: int = 15,
          department_id: int | None = None) -> list[dict]:
    """[{id, name, head, rows: [{employee_id, full_name, position, shtat, cells, locked, ik, ikm}]}]"""
    days_in_month = monthrange(year, month)[1]
    day_to = min(day_to, days_in_month)
    days = list(range(day_from, day_to + 1))
    today = datetime.now(TZ_UZ).date()
    d_from, d_to = f"{year:04d}-{month:02d}-{day_from:02d}", f"{year:04d}-{month:02d}-{day_to:02d}"

    # is_active bu yerda filtr emas: ta'til/bolnichniy/safar paytida is_active=False bo'ladi,
    # lekin xodim tabelda turishi va kunlariga М/Т, К ... qo'yilishi kerak
    q = db.query(models.Employee).filter(models.Employee.department_id.isnot(None),
                                         models.Employee.role.notin_(_EXCLUDED_ROLES))
    if department_id is not None:
        q = q.filter(models.Employee.department_id == department_id)
    employees = q.all()
    emp_ids = [e.id for e in employees]

    holidays = month_holidays(db, year, month)
    statuses = status_code_map(db, employees, d_from, d_to)
    att = {(a.employee_id, a.date) for a in db.query(models.Attendance.employee_id, models.Attendance.date)
           .filter(models.Attendance.employee_id.in_(emp_ids), models.Attendance.date >= d_from,
                   models.Attendance.date <= d_to).all()} if emp_ids else set()
    overrides = {(o.employee_id, o.date): o.code for o in db.query(models.TabelOverride)
                 .filter(models.TabelOverride.employee_id.in_(emp_ids), models.TabelOverride.date >= d_from,
                         models.TabelOverride.date <= d_to).all()} if emp_ids else {}
    manual = {(r.employee_id, r.day): r.code for r in db.query(models.TabelRecord)
              .filter(models.TabelRecord.employee_id.in_(emp_ids), models.TabelRecord.year == year,
                      models.TabelRecord.month == month).all()} if emp_ids else {}

    def cell(emp, day: int) -> tuple[str, bool]:
        """(kod, qulflangan) — qulflangan kunni bo'lim boshlig'i o'zgartira olmaydi."""
        d = date(year, month, day)
        iso = d.isoformat()
        if (emp.id, iso) in overrides:
            code = overrides[(emp.id, iso)]
            return (_work_code(emp) if code == "8" else DISPLAY.get(code, code)), True
        if d.weekday() >= 5:
            return "X", True
        if day in holidays:
            return "Б", True
        if (emp.id, iso) in statuses:
            return DISPLAY[statuses[(emp.id, iso)]], True
        if emp.status == S.dekret:
            return "Д", True
        code = _work_code(emp) if (emp.status == S.shafyor_farrosh or (emp.id, iso) in att or d > today) else ""
        m = manual.get((emp.id, day))
        if m and m not in ("8", "4") and m != code:
            code = m   # bo'lim boshlig'i qo'lda belgilagan (masalan, К yoki У/Т)
        return code, False

    by_dept: dict[int, list] = {}
    for e in employees:
        by_dept.setdefault(e.department_id, []).append(e)
    depts = (db.query(models.Department).filter(models.Department.id.in_(list(by_dept)))
             .order_by(models.Department.order_num, models.Department.id).all()) if by_dept else []

    out = []
    for dep in depts:
        staff = by_dept[dep.id]
        head = next((e for e in staff if e.has_role(*_HEAD_ROLES)), None)
        # Bo'lim boshlig'i — birinchi qatorda, qolganlar — familiya bo'yicha
        staff.sort(key=lambda e: (e is not head, e.full_name))
        rows = []
        for e in staff:
            cells, locked = {}, []
            for day in days:
                code, lock = cell(e, day)
                cells[str(day)] = code
                if lock:
                    locked.append(day)
            vals = cells.values()
            rows.append({
                "employee_id": e.id, "full_name": e.full_name, "position": e.position or "",
                "work_rate": e.work_rate or 1, "shtat": f"{fmt_rate(e.work_rate or 1)} st.",
                "is_head": e is head, "cells": cells, "locked": locked,
                "ik": sum(1 for v in vals if v in WORK_CODES),
                "ikm": sum(1 for v in vals if v in ABSENT_CODES),
            })
        out.append({"id": dep.id, "name": dep.name, "head": head.full_name if head else None,
                    "head_position": head.position if head else None, "rows": rows})
    return out


def fmt_rate(r: float) -> str:
    return (f"{r:.2f}".rstrip("0").rstrip(".")).replace(".", ",")


# ── Word (DOCX) ──────────────────────────────────────────────────────────────

NAVY, HEAD_BG, DAY_BG, BORDER, MUTED, WKND_BG = "17365D", "245A8D", "EAF2F8", "C7D6E5", "526276", "E9EAEC"
CODE_STYLE = {   # kod: (fon, matn rangi)
    "4": ("F0F9FF", "0369A1"), "X": (WKND_BG, "666666"), "Б": ("FEFCE8", "854D0E"), "К": ("FEF2F2", "991B1B"),
    "У/Т": ("EFF6FF", "1D4ED8"), "М/Т": ("FAF5FF", "6D28D9"), "Х/С": ("ECFDF5", "047857"),
    "Б/С": ("FFF7ED", "C2410C"), "Д": ("FDF2F8", "BE185D"),
}
MONTHS = ["", "yanvar", "fevral", "mart", "aprel", "may", "iyun", "iyul", "avgust", "sentabr", "oktabr", "noyabr", "dekabr"]
DIRECTOR = "R.Kuchkarov"


def build_docx(depts: list[dict], year: int, month: int, day_from: int, day_to: int) -> bytes:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, Cm, RGBColor

    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
    sec.left_margin = sec.right_margin = Cm(1.27)
    sec.top_margin = sec.bottom_margin = Cm(1.4)
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(10)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    st.paragraph_format.space_after = Pt(0)
    usable = 29.7 - 2 * 1.27   # sm

    def para(text="", size=10, bold=False, color=NAVY, align=None, space_after=0, container=None):
        p = (container or doc).add_paragraph()
        if align is not None:
            p.alignment = align
        p.paragraph_format.space_after = Pt(space_after)
        if text:
            r = p.add_run(text)
            r.bold, r.font.size, r.font.color.rgb = bold, Pt(size), RGBColor.from_string(color)
        return p

    def shade(cell, fill):
        tcPr = cell._tc.get_or_add_tcPr()
        sh = OxmlElement("w:shd")
        sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), fill)
        tcPr.append(sh)

    def borders(tbl, color=BORDER, size=4, none=False):
        tblPr = tbl._tbl.tblPr
        b = OxmlElement("w:tblBorders")
        for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
            e = OxmlElement(f"w:{side}")
            e.set(qn("w:val"), "nil" if none else "single")
            if not none:
                e.set(qn("w:sz"), str(size)); e.set(qn("w:color"), color)
            b.append(e)
        tblPr.append(b)

    def fixed(tbl, widths_cm):
        tbl.autofit = False
        lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tbl._tbl.tblPr.append(lay)
        for row in tbl.rows:
            for i, w in enumerate(widths_cm):
                row.cells[i].width = Cm(w)

    def put(cell, text, size=10, bold=False, color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER, fill=None):
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = align
        p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(2)
        if text:
            r = p.add_run(text)
            r.bold, r.font.size, r.font.color.rgb = bold, Pt(size), RGBColor.from_string(color)
        if fill:
            shade(cell, fill)

    days = list(range(day_from, day_to + 1))
    title = ("DAVOMAT JADVALI — YARIM OYLIK (1–15 KUN)" if (day_from, day_to) == (1, 15)
             else f"DAVOMAT JADVALI ({day_from}–{day_to} KUN)")
    mm = f"{month:02d}"

    for di, dep in enumerate(depts):
        if di:
            para().add_run().add_break(WD_BREAK.PAGE)

        # «TASDIQLAYMAN»
        t = doc.add_table(rows=1, cols=2); borders(t, none=True); fixed(t, [usable * 0.58, usable * 0.42])
        c = t.rows[0].cells[1]
        p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = p.add_run("«TASDIQLAYMAN»"); r.bold, r.font.size, r.font.color.rgb = True, Pt(13), RGBColor.from_string(NAVY)
        for line, b in (("Texnik me'yorlash va standartlashtirish", False), ("ilmiy-tadqiqot instituti direktori v.v.b.", False),
                        (DIRECTOR, True), (f"«___»__________{year} y.", False)):
            para(line, 9.5, b, NAVY if b else MUTED, WD_ALIGN_PARAGRAPH.RIGHT, container=c)

        para(title, 17, True, NAVY, WD_ALIGN_PARAGRAPH.CENTER).paragraph_format.space_before = Pt(6)
        para(f"{dep['name']}  •  {day_from:02d}.{mm}.{year} — {day_to:02d}.{mm}.{year}", 10.5, False, MUTED,
             WD_ALIGN_PARAGRAPH.CENTER, space_after=6)

        # Jadval: № | F.I.Sh. | Lavozim | Shtat | kunlar... | I.K. | I.Km.
        fixed_w = [0.8, 3.6, 3.2, 1.3]
        sum_w = [1.1, 1.4]
        day_w = max(0.55, (usable - sum(fixed_w) - sum(sum_w)) / len(days))
        widths = fixed_w + [day_w] * len(days) + sum_w
        ncol = len(widths)
        tbl = doc.add_table(rows=2 + len(dep["rows"]), cols=ncol)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        borders(tbl); fixed(tbl, widths)
        h0, h1 = tbl.rows[0].cells, tbl.rows[1].cells
        for i, txt in ((0, "№"), (1, "F.I.Sh."), (2, "Lavozim"), (3, "Shtat"), (ncol - 2, "I.K."), (ncol - 1, "I.Km.")):
            m = h0[i].merge(h1[i])
            put(m, txt, 9.5, True, "FFFFFF", fill=HEAD_BG)
        kun = h0[4].merge(h0[4 + len(days) - 1])
        put(kun, "KUNLAR", 9.5, True, "FFFFFF", fill=HEAD_BG)
        for j, d in enumerate(days):
            put(h1[4 + j], str(d), 9, True, NAVY, fill=DAY_BG)
        for row_el in (tbl.rows[0], tbl.rows[1]):   # sarlavha har sahifada takrorlanadi
            trPr = row_el._tr.get_or_add_trPr()
            th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true"); trPr.append(th)

        for ri, row in enumerate(dep["rows"]):
            cs = tbl.rows[2 + ri].cells
            put(cs[0], str(ri + 1), 9.5, True)
            put(cs[1], row["full_name"], 9.5, row["is_head"], align=WD_ALIGN_PARAGRAPH.LEFT)
            put(cs[2], row["position"], 8.5)
            put(cs[3], row["shtat"], 9)
            for j, d in enumerate(days):
                code = row["cells"].get(str(d), "")
                bg, fg = CODE_STYLE.get(code, ("FFFFFF", NAVY))
                put(cs[4 + j], code, 9, False, fg, fill=bg if bg != "FFFFFF" else None)
            put(cs[ncol - 2], str(row["ik"] or ""), 10.5, True)
            put(cs[ncol - 1], str(row["ikm"] or ""), 10.5, True, "C0392B")

        # Shartli belgilar | Bo'lim boshlig'i imzosi
        para(space_after=10)
        foot = doc.add_table(rows=1, cols=2); borders(foot, none=True); fixed(foot, [usable * 0.5, usable * 0.5])
        lc, rc = foot.rows[0].cells
        nrows = (len(LEGEND) + 1) // 2
        leg = lc.add_table(rows=nrows, cols=4); borders(leg, none=True)
        fixed(leg, [1.0, 4.2, 1.0, 4.2])
        for k, (code, label) in enumerate(LEGEND):
            rr, cc = k % nrows, (k // nrows) * 2
            bg, fg = CODE_STYLE.get(code, ("FFFFFF", NAVY))
            put(leg.rows[rr].cells[cc], code, 9, True, fg, fill=bg)
            put(leg.rows[rr].cells[cc + 1], label, 9, False, MUTED, WD_ALIGN_PARAGRAPH.LEFT)
        lc._tc.append(lc.paragraphs[0]._p)   # katak oxiri paragraf bo'lishi shart (Word talabi)
        para("BO'LIM BOSHLIG'I", 11.5, True, NAVY, WD_ALIGN_PARAGRAPH.LEFT, space_after=4, container=rc)
        if dep.get("head"):
            para(dep["head"], 10.5, True, NAVY, WD_ALIGN_PARAGRAPH.LEFT, space_after=4, container=rc)
        para("Imzo: ______________________", 10, False, MUTED, WD_ALIGN_PARAGRAPH.LEFT, space_after=4, container=rc)
        para(f"Sana: «___»__________{year} y.", 10, False, MUTED, WD_ALIGN_PARAGRAPH.LEFT, container=rc)
        # Bo'sh birinchi paragraflarni olib tashlash (katak boshida ortiqcha qator bo'lmasin)
        first = rc.paragraphs[0]
        if not first.text and len(rc.paragraphs) > 1:
            first._p.getparent().remove(first._p)

    if not depts:
        para("Ma'lumot yo'q", 12, True, MUTED, WD_ALIGN_PARAGRAPH.CENTER)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
