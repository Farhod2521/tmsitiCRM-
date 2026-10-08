"""Energoaudit hisoboti — "Хисобот Уй-жойлар шаблон (2).docx" asosida.

Shablondagi sariq (highlight) joylar formadan to'ldiriladi, hisob-kitoblar
(R₀, Q, ulushlar, "мос/мос эмас", xulosa va tavsiyalardagi qiymatlar) shu yerda
bajariladi. Elementlar shablondagi tartib raqami bo'yicha olinadi (shablon
o'zgarmas — app/templates/energoaudit_uyjoy.docx). Shablon almashtirilsa,
quyidagi indekslar ham qayta tekshirilishi kerak."""
import base64
import copy
import io
import os
from datetime import datetime
from typing import Any, Optional

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm

TEMPLATE_PATH = os.path.join(os.path.dirname(__file__), "templates", "energoaudit_uyjoy.docx")

# Deraza/eshik maydonlari soni bo'yicha hisoblanadi (bitta element o'rtacha maydoni, m²)
DERAZA_K, ESHIK_K, YOLAK_K = 2.2, 1.9, 1.5

# Suratlar: muqova + hujjatdagi 3 ta rasm (sarlavhasi shablonda)
PHOTO_SLOTS = {
    "muqova": "Muqova — bino surati",
    "rasm_1": "1-расм. Девор конструкциясининг қатламлари",
    "rasm_2": "2-расм. Томёпма конструкциясининг қатламлари",
    "rasm_3": "3-расм. Иссиқлик йўқотишларининг тепловизион тасвири",
}

DEFAULT_DATA: dict[str, Any] = {
    "viloyat": "",
    "tuman": "",
    "bino_nomi": "",
    "qurilgan_yil": 1973,
    "umumiy_maydon": 2906,
    "qavatlar_soni": 3,
    "fasad_izolyatsiya": "мавжуд эмас",
    "isitish_turi": "марказий",
    "issiq_suv": "маҳаллий",
    "resurslar": "электр энергияси ва иссиқлик энергиясини",
    "issiqlik_manbai": "марказий иссиқлик таъминоти",
    "uzunlik": 38.7,
    "kenglik": 24.4,
    "balandlik": 10.8,
    "pastki_konstruksiya": "ертўла ораёпмаси",
    "deraza_soni": 88,
    "yolak_deraza_soni": 10,
    "eshik_soni": 6,
    "devor_maydon": None,            # bo'sh bo'lsa: 2·(L+W)·H − derazalar − eshiklar
    "tom_maydon": 648.63,
    "pol_maydon": None,              # bo'sh bo'lsa: L·W
    "tom_turi": "чордоқли икки нишабли",
    "tom_tavsif": "Томёпмаси темир-бетон плиталардан бажарилган бўлиб, иссиқлик изоляцияси сифатида керамзит тўшалган.",
    "devor_qatlamlar": [
        {"nomi": "Ички сувоқ", "qalinlik": 15, "lambda": 0.93},
        {"nomi": "Пишиқ ғишт", "qalinlik": 380, "lambda": 0.52},
        {"nomi": "Ташқи сувоқ", "qalinlik": 15, "lambda": 0.93},
    ],
    "tom_qatlamlar": [
        {"nomi": "Темир-бетон плита", "qalinlik": 220, "lambda": 2.04},
        {"nomi": "Цемент-қум стяжкаси", "qalinlik": 150, "lambda": 0.93},
    ],
    "pol_qatlamlar": [],             # bo'sh — jadvalda "Мавжуд эмас"
    "alfa_ichki": 8.7, "alfa_devor": 23, "alfa_tom": 12,
    "alfa_pol_ichki": 8.7, "alfa_pol": 6,
    "pol_n": 0.6,                    # ҚМҚ 2.01.04-18: ташқи ҳавога нисбатан коэффициент (иситилмайдиган ертўла устидаги пол)
    "deraza_R": 0.62, "eshik_R": 0.56,
    "t_ichki": 20, "t_tashqi": -15,
    "bino_toifa": "Кўп қаватли уй-жойлар",
    "dd": 2877,
    "norma_toifa": "3 қаватдан юқори турар-жой бинолари ва даволаш-профилактика муассасалари",
    "norma_dd": "2000-3000",
    "norma_devor": 2.2, "norma_tom": 3.0, "norma_pol": 2.8, "norma_deraza": 0.53, "norma_eshik": 0.53, "norma_fonar": 0.31,
}

# Shablon elementlari (body ichidagi tartib raqami)
E = dict(
    cover_title=17, cover_tbl=19, cover_year=20,
    intro=32, yil=34, maydon=35, qavat=36, fasad=37, isitish=38, issiq_suv=39, resurs=41,
    hajm=44, pastki=45, devor_tavsif=46, deraza_matn=47, deraza_tbl=50, tom_matn=52,
    rasm1_pic=69, rasm1_cap=70,
    devor_q=(72, 73, 74), devor_r=(77, 78, 79), devor_rk=81, devor_r0=83,
    tom_sar=85, tom_q=(86, 87, 88), tom_r=(90, 91, 92), tom_rk=94, tom_r0=96,
    rasm2_pic=97, rasm2_cap=98, pol_after=99,
    deraza_R=100, norma_tbl=104,
    q_devor=(117, 118), q_tom=(119, 121), q_deraza=(122, 123), q_eshik=(124, 125),
    q_jami=126, ulush=(128, 129, 130, 131),
    rasm3_pic=132, rasm3_cap=133,
    xulosa1=135, xulosa2=136, t_devor=140, t_tom=143, t_deraza=145, t_eshik=147,
)


# ── Hisob-kitob ──────────────────────────────────────────────────────────────

def _num(v, default=0.0) -> float:
    try:
        if v is None or v == "":
            return float(default)
        return float(str(v).replace(",", ".").replace(" ", ""))
    except (TypeError, ValueError):
        return float(default)


def fmt(v: Optional[float], nd: int = 2) -> str:
    """O'zbekcha son: vergul bilan, keraksiz nollarsiz (0,921; 41056,5; 2906)."""
    if v is None:
        return "—"
    s = f"{v:.{nd}f}"
    if "." in s:
        s = s.rstrip("0").rstrip(".")
    return s.replace(".", ",")


def nfmt(v: float) -> str:
    """Me'yoriy qiymat: kamida bitta kasr bilan (3,0; 2,2; 0,53)."""
    return fmt(v, 2) if round(v, 1) != v else f"{v:.1f}".replace(".", ",")


def merged(data: dict) -> dict:
    d = dict(DEFAULT_DATA)
    d.update({k: v for k, v in (data or {}).items() if v is not None or k in ("devor_maydon", "pol_maydon")})
    return d


def _layers(raw) -> list[dict]:
    out = []
    for l in raw or []:
        t, lam = _num(l.get("qalinlik")), _num(l.get("lambda"))
        if t <= 0 and lam <= 0 and not (l.get("nomi") or "").strip():
            continue   # to'liq bo'sh qator
        r = round(t / 1000 / lam, 3) if t > 0 and lam > 0 else 0.0
        out.append({"nomi": (l.get("nomi") or "").strip() or "Қатлам", "qalinlik": t, "lambda": lam, "R": r})
    return out


def _r0(layers, a_in, a_out) -> tuple[float, float]:
    rk = round(sum(l["R"] for l in layers), 3)
    r0 = round(1 / a_in + rk + 1 / a_out, 3) if a_in and a_out else rk
    return rk, r0


def calculate(data: dict) -> dict:
    d = merged(data)
    ai = _num(d["alfa_ichki"], 8.7)
    dt = _num(d["t_ichki"], 20) - _num(d["t_tashqi"], -15)

    wall_layers, roof_layers, floor_layers = _layers(d["devor_qatlamlar"]), _layers(d["tom_qatlamlar"]), _layers(d.get("pol_qatlamlar"))
    wall_rk, wall_r0 = _r0(wall_layers, ai, _num(d["alfa_devor"], 23))
    roof_rk, roof_r0 = _r0(roof_layers, ai, _num(d["alfa_tom"], 12))
    floor_rk, floor_r0 = _r0(floor_layers, _num(d["alfa_pol_ichki"], 8.7), _num(d["alfa_pol"], 6))
    has_floor = bool(floor_layers)

    L, W, H = _num(d["uzunlik"]), _num(d["kenglik"]), _num(d["balandlik"])
    hajm = round(L * W * H, 2)
    n_win, n_lobby, n_door = int(_num(d["deraza_soni"])), int(_num(d["yolak_deraza_soni"])), int(_num(d["eshik_soni"]))
    a_win, a_lobby, a_door = round(n_win * DERAZA_K, 2), round(n_lobby * YOLAK_K, 2), round(n_door * ESHIK_K, 2)
    win_n, win_a = n_win + n_lobby, round(a_win + a_lobby, 2)

    wall_auto = round(max(0.0, 2 * (L + W) * H - win_a - a_door), 2)
    wall_a = round(_num(d["devor_maydon"]), 2) if d.get("devor_maydon") not in (None, "") else wall_auto
    roof_a = round(_num(d["tom_maydon"]), 2)
    floor_auto = round(L * W, 2)
    floor_a = round(_num(d["pol_maydon"]), 2) if d.get("pol_maydon") not in (None, "") else floor_auto
    rw, rd = _num(d["deraza_R"], 0.62), _num(d["eshik_R"], 0.56)

    def q(a, r):
        return round(a * dt / r, 1) if r > 0 else 0.0

    losses = [
        {"key": "tom", "nomi": "Том", "A": roof_a, "R": roof_r0, "Q": q(roof_a, roof_r0)},
        {"key": "devor", "nomi": "Деворлар", "A": wall_a, "R": wall_r0, "Q": q(wall_a, wall_r0)},
        {"key": "deraza", "nomi": "Деразалар", "A": win_a, "R": rw, "Q": q(win_a, rw)},
        {"key": "eshik", "nomi": "Эшиклар", "A": a_door, "R": rd, "Q": q(a_door, rd)},
    ]
    if has_floor:
        n_pol = _num(d.get("pol_n"), 0.6)
        losses.append({"key": "pol", "nomi": "Пол", "A": floor_a, "R": floor_r0, "n": n_pol,
                       "Q": round(n_pol * q(floor_a, floor_r0), 1)})
    total_w = sum(x["Q"] for x in losses)
    for x in losses:
        x["kw"] = round(x["Q"] / 1000, 2)
        x["ulush"] = round(x["Q"] / total_w * 100, 1) if total_w else 0.0

    n = {k: _num(d[f"norma_{k}"]) for k in ("devor", "tom", "pol", "deraza", "eshik")}
    compare = {
        "devor": wall_r0 >= n["devor"],
        "tom": roof_r0 >= n["tom"],
        "pol": floor_r0 >= n["pol"] if has_floor else None,
        "deraza": rw >= n["deraza"],
        "eshik": rd >= n["eshik"],
    }
    return {
        "dt": dt, "hajm": hajm,
        "devor": {"qatlamlar": wall_layers, "Rk": wall_rk, "R0": wall_r0},
        "tom": {"qatlamlar": roof_layers, "Rk": roof_rk, "R0": roof_r0},
        "pol": {"qatlamlar": floor_layers, "Rk": floor_rk, "R0": floor_r0 if has_floor else None},
        "deraza_soni": win_n, "deraza_maydon": win_a,
        "deraza_maydon_asosiy": a_win, "yolak_deraza_maydon": a_lobby,
        "eshik_soni": n_door, "eshik_maydon": a_door,
        "devor_maydon": wall_a, "devor_maydon_auto": wall_auto, "tom_maydon": roof_a,
        "pol_maydon": floor_a, "pol_maydon_auto": floor_auto,
        "yoqotishlar": losses, "jami_kw": round(total_w / 1000, 1),
        "taqqoslash": compare,
        "koef": {"deraza": DERAZA_K, "eshik": ESHIK_K, "yolak": YOLAK_K},
    }


# ── DOCX yordamchilari ───────────────────────────────────────────────────────

Seg = tuple[str, Optional[str]]   # (matn, None | "subscript" | "superscript")


def _base_rpr(p):
    for r in p.iter(qn("w:r")):
        if any((t.text or "").strip() for t in r.iter(qn("w:t"))):
            rpr = r.find(qn("w:rPr"))
            if rpr is None:
                return None
            rpr = copy.deepcopy(rpr)
            for tag in ("w:highlight", "w:vertAlign", "w:shd"):
                for e in rpr.findall(qn(tag)):
                    rpr.remove(e)
            return rpr
    return None


def set_text(p, content) -> None:
    """Paragraf matnini to'liq almashtiradi (birinchi matnli run formati, sariq fonsiz).
    content — str yoki [(matn, vertAlign), ...]."""
    segs: list[Seg] = [(content, None)] if isinstance(content, str) else list(content)
    rpr = _base_rpr(p)
    for child in list(p):
        if child.tag in (qn("w:r"), qn("w:hyperlink"), qn("w:smartTag"), qn("w:proofErr"),
                         qn("w:bookmarkStart"), qn("w:bookmarkEnd")):
            p.remove(child)
    for text, valign in segs:
        if not text:
            continue
        r = OxmlElement("w:r")
        if rpr is not None or valign:
            rp = copy.deepcopy(rpr) if rpr is not None else OxmlElement("w:rPr")
            if valign:
                va = OxmlElement("w:vertAlign"); va.set(qn("w:val"), valign); rp.append(va)
            r.append(rp)
        t = OxmlElement("w:t"); t.set(qn("xml:space"), "preserve"); t.text = text
        r.append(t)
        p.append(r)


def _cell(tbl, row, col):
    tr = tbl.findall(qn("w:tr"))[row]
    return tr.findall(qn("w:tc"))[col]


def set_cell(tbl, row, col, content) -> None:
    tc = _cell(tbl, row, col)
    ps = tc.findall(qn("w:p"))
    set_text(ps[0], content)
    for extra in ps[1:]:
        tc.remove(extra)


def replace_lines(protos: list, lines: list) -> None:
    """Bir xil turdagi qatorlar guruhini (masalan, devor qatlamlari) yangi ro'yxat bilan almashtiradi."""
    last = protos[-1]
    for i, content in enumerate(lines):
        if i < len(protos):
            set_text(protos[i], content)
            last = protos[i]
        else:
            new = copy.deepcopy(protos[-1])
            last.addnext(new)
            set_text(new, content)
            last = new
    for extra in protos[len(lines):]:
        extra.getparent().remove(extra)


def insert_after(anchor, proto, contents: list):
    """proto paragrafidan nusxa olib, anchor'dan keyin ketma-ket qo'yadi. Oxirgisini qaytaradi."""
    last = anchor
    for content in contents:
        new = copy.deepcopy(proto)
        set_text(new, content)
        last.addnext(new)
        last = new
    return last


def remove(*elements) -> None:
    for e in elements:
        if e is not None and e.getparent() is not None:
            e.getparent().remove(e)


def _image(data_url: str) -> Optional[tuple[io.BytesIO, float]]:
    """(oqim, eni/bo'yi nisbati) — buzilgan bo'lsa None."""
    try:
        raw = base64.b64decode(data_url.split(",", 1)[1] if data_url.startswith("data:") else data_url)
    except Exception:
        return None
    ratio = 4 / 3
    try:
        from PIL import Image
        with Image.open(io.BytesIO(raw)) as im:
            ratio = im.width / im.height if im.height else ratio
    except Exception:
        pass
    return io.BytesIO(raw), ratio


def _fit(ratio: float, max_w: float, max_h: float) -> dict:
    """Rasmni nisbatini saqlab max_w × max_h (sm) ichiga sig'diradi."""
    return {"width": Cm(max_w)} if max_w / ratio <= max_h else {"height": Cm(max_h)}


def _put_picture(document, p_el, data_url: str, max_w: float, max_h: float) -> bool:
    from docx.text.paragraph import Paragraph
    img = _image(data_url or "")
    if img is None:
        return False
    stream, ratio = img
    para = Paragraph(p_el, document._body)
    for r in list(p_el.findall(qn("w:r"))):
        p_el.remove(r)
    pPr = p_el.find(qn("w:pPr"))
    if pPr is None:
        pPr = OxmlElement("w:pPr"); p_el.insert(0, pPr)
    for jc in pPr.findall(qn("w:jc")):
        pPr.remove(jc)
    jc = OxmlElement("w:jc"); jc.set(qn("w:val"), "center"); pPr.append(jc)
    try:
        para.add_run().add_picture(stream, **_fit(ratio, max_w, max_h))
        return True
    except Exception:
        return False


def _pct_below(r: float, norm: float) -> int:
    return round((norm - r) / norm * 100) if norm else 0


# ── DOCX yaratish ────────────────────────────────────────────────────────────

def build_docx(data: dict, photos: Optional[dict] = None, title: str = "") -> bytes:
    d = merged(data)
    c = calculate(d)
    photos = photos or {}
    document = docx.Document(TEMPLATE_PATH)
    els = list(document.element.body.iterchildren())
    P = lambda key: els[E[key]]
    dt = fmt(c["dt"], 1)
    U = "м²·°С/Вт"

    # ── Muqova
    obyekt = (title or d.get("bino_nomi") or "").strip()
    set_text(P("cover_title"), f"{obyekt} ЭНЕРГОАУДИТИ ҲИСОБОТИ" if obyekt else "ЭНЕРГОАУДИТ ҲИСОБОТИ")
    set_text(P("cover_year"), f"Тошкент {datetime.now().year}")
    cover_tbl = P("cover_tbl")
    if photos.get("muqova"):
        _put_picture(document, _cell(cover_tbl, 0, 0).find(qn("w:p")), photos["muqova"], 14.6, 10.0)
    else:
        remove(cover_tbl)   # suratsiz bo'sh ramka qolmasin

    # ── 2.1–2.2 Umumiy ma'lumot
    joy = ", ".join(x for x in (str(d.get("viloyat") or "").strip(), str(d.get("tuman") or "").strip()) if x)
    bino = str(d.get("bino_nomi") or "").strip()
    obj = ", ".join(x for x in (joy, bino) if x) or "бинода"
    set_text(P("intro"), f"Энергоаудит {obj} Қурилиш ва уй-жой коммунал хўжалиги вазирлиги ҳузуридаги Техник меъёрлаш "
                         "ва стандартлаштириш илмий-тадқиқот институти томонидан ўтказилди.")
    set_text(P("yil"), f"Қурилган йил – {fmt(_num(d['qurilgan_yil']), 0)} йил")
    set_text(P("maydon"), f"Умумий майдони – {fmt(_num(d['umumiy_maydon']))} м²")
    set_text(P("qavat"), f"Қаватлар сони – {fmt(_num(d['qavatlar_soni']), 0)}")
    set_text(P("fasad"), f"Фасад қисмида иссиқлик ҳимоя қопламаси – {d['fasad_izolyatsiya']}")
    set_text(P("isitish"), f"Иситиш тизими тури – {d['isitish_turi']}")
    set_text(P("issiq_suv"), f"Иссиқ сув таъминоти тури – {d['issiq_suv']}")
    set_text(P("resurs"), f"Биноларда фаолияти давомида ёқилғи-энергетика ресурсларидан {d['resurslar']} истеъмол "
                          f"қилиниши аниқланди. Бунда иссиқлик энергияси {d['issiqlik_manbai']} орқали таъминланади.")

    # ── 2.3 Hajm-reja yechimlari
    set_text(P("hajm"), f"Бино {fmt(_num(d['qavatlar_soni']), 0)} қаватдан иборат бўлиб, режада {fmt(_num(d['uzunlik']))} × "
                        f"{fmt(_num(d['kenglik']))} м ўлчамга эга. Бинонинг умумий майдони {fmt(_num(d['umumiy_maydon']))} м², "
                        f"умумий ҳажми {fmt(c['hajm'], 1)} м³.")
    set_text(P("pastki"), f"Ташқи тўсувчи конструкциялар ташқи деворлар, дераза ва эшиклар, {d['pastki_konstruksiya']} ҳамда "
                          "том конструкцияларидан ташкил топган.")
    wl = c["devor"]["qatlamlar"]
    main = max(wl, key=lambda l: l["qalinlik"]) if wl else None
    others = [l for l in wl if l is not main]
    wall_desc = f"{fmt(main['qalinlik'], 0)} мм қалинликдаги {main['nomi'].lower()}" if main else "—"
    if others:
        wall_desc += ", " + ", ".join(f"{fmt(l['qalinlik'], 0)} мм {l['nomi'].lower()}" for l in others)
    set_text(P("devor_tavsif"), f"Бинонинг ташқи деворлари {wall_desc}дан иборат.")
    set_text(P("deraza_matn"), [
        (f"Бино ҳар бир томондаги фасадида умумий майдони {fmt(c['deraza_maydon'])} м", None), ("2", "superscript"),
        (f" ни ташкил қиладиган {c['deraza_soni']} та деразалар ва умумий майдони {fmt(c['eshik_maydon'])} м", None),
        ("2", "superscript"), (f" ни ташкил қиладиган {c['eshik_soni']} та эшиклар мавжуд.", None)])
    t = P("deraza_tbl")
    set_cell(t, 1, 1, fmt(_num(d["deraza_soni"]), 0)); set_cell(t, 2, 1, fmt(c["deraza_maydon_asosiy"]))
    set_cell(t, 1, 2, fmt(_num(d["eshik_soni"]), 0)); set_cell(t, 2, 2, fmt(c["eshik_maydon"]))
    set_cell(t, 1, 3, fmt(_num(d["yolak_deraza_soni"]), 0)); set_cell(t, 2, 3, fmt(c["yolak_deraza_maydon"]))
    set_text(P("tom_matn"), f"Бинонинг том қисми {d['tom_turi']} сифатида ташкил қилинган. {d['tom_tavsif']}".strip())

    # ── 2.4.1 R₀ — devor, tom, pol
    def layer_lines(layers):
        return [f"{l['nomi']} – {fmt(l['qalinlik'], 0)} мм" for l in layers] or ["—"]

    def r_lines(layers):
        return [[("R", None), (str(i + 1), "subscript"),
                 (f" ({l['nomi'].lower()}) = {fmt(l['qalinlik'] / 1000, 3)}/{fmt(l['lambda'], 3)} = {fmt(l['R'], 3)}", None)]
                for i, l in enumerate(layers)] or [["—"]]

    def rk_line(part):
        terms = " + ".join(fmt(l["R"], 3) for l in part["qatlamlar"]) or "0"
        return [("R", None), ("k", "subscript"), (f" = {terms} = {fmt(part['Rk'], 3)}", None)]

    def r0_line(part, a_in, a_out):
        return [("R", None), ("o", "subscript"),
                (f" = (1/{fmt(_num(a_in), 1)}) + {fmt(part['Rk'], 3)} + (1/{fmt(_num(a_out), 1)}) = {fmt(part['R0'], 3)}", None)]

    replace_lines([els[i] for i in E["devor_q"]], layer_lines(wl))
    replace_lines([els[i] for i in E["devor_r"]], r_lines(wl))
    set_text(P("devor_rk"), rk_line(c["devor"]))
    set_text(P("devor_r0"), r0_line(c["devor"], d["alfa_ichki"], d["alfa_devor"]))
    rl = c["tom"]["qatlamlar"]
    replace_lines([els[i] for i in E["tom_q"]], layer_lines(rl))
    replace_lines([els[i] for i in E["tom_r"]], r_lines(rl))
    set_text(P("tom_rk"), rk_line(c["tom"]))
    set_text(P("tom_r0"), r0_line(c["tom"], d["alfa_ichki"], d["alfa_tom"]))

    # Pol — shablonda yo'q: tom bo'limi namunasida, 2-rasmdan keyin qo'shiladi
    if c["pol"]["R0"] is not None:
        pl = c["pol"]["qatlamlar"]
        last = insert_after(P("pol_after"), P("tom_sar"), [f"Полнинг ({d['pastki_konstruksiya']}) иссиқлик узатилишига қаршилиги"])
        last = insert_after(last, els[E["tom_q"][0]], layer_lines(pl))
        last = insert_after(last, els[E["tom_r"][0]], r_lines(pl))
        last = insert_after(last, P("tom_rk"), [rk_line(c["pol"])])
        insert_after(last, P("tom_r0"), [r0_line(c["pol"], d["alfa_pol_ichki"], d["alfa_pol"])])

    set_text(P("deraza_R"), f"Дераза ва эшикнинг иссиқлик узатишга қаршилиги мос равишда {fmt(_num(d['deraza_R']), 3)} ва "
                            f"{fmt(_num(d['eshik_R']), 3)} деб қабул қилинади.")

    # 2-jadval — solishtirma
    t = P("norma_tbl")
    for col, val in enumerate([d["norma_toifa"], d["norma_dd"], nfmt(_num(d["norma_devor"])), nfmt(_num(d["norma_tom"])),
                               nfmt(_num(d["norma_pol"])), nfmt(_num(d["norma_deraza"])), nfmt(_num(d["norma_fonar"]))]):
        set_cell(t, 2, col, str(val))
    pol_r0 = c["pol"]["R0"]
    for col, val in enumerate([d["bino_toifa"], fmt(_num(d["dd"]), 0), fmt(c["devor"]["R0"], 3), fmt(c["tom"]["R0"], 3),
                               "Мавжуд эмас" if pol_r0 is None else fmt(pol_r0, 3),
                               fmt(_num(d["deraza_R"]), 3), "Мавжуд эмас"]):
        set_cell(t, 3, col, val)
    mos = lambda ok: "" if ok is None else ("мос" if ok else "мос эмас")
    cmp_ = c["taqqoslash"]
    set_cell(t, 4, 0, "Таққослаш")
    for col, ok in ((2, cmp_["devor"]), (3, cmp_["tom"]), (4, cmp_["pol"]), (5, cmp_["deraza"])):
        set_cell(t, 4, col, mos(ok))

    # ── 2.4.2 Issiqlik yo'qotishlari
    by = {x["key"]: x for x in c["yoqotishlar"]}
    q_line = lambda x: (f"Q = {fmt(x['n'], 2) + ' × ' if 'n' in x else ''}{fmt(x['A'])} × {dt} / {fmt(x['R'], 3)}"
                        f" = {fmt(x['Q'], 1)} Вт")
    for key, label in (("devor", "Деворлардан иссиқлик йўқотилиши"), ("tom", "Том орқали иссиқлик йўқотилиши"),
                       ("deraza", "Деразалар орқали иссиқлик йўқотилиши"), ("eshik", "Эшиклар орқали иссиқлик йўқотилиши")):
        li, qi = E[f"q_{key}"]
        set_text(els[li], label); set_text(els[qi], q_line(by[key]))
    if "pol" in by:
        li, qi = E["q_eshik"]
        last = insert_after(els[qi], els[li], ["Пол орқали иссиқлик йўқотилиши"])
        insert_after(last, els[qi], [q_line(by["pol"])])
    set_text(P("q_jami"), f"Тўсувчи конструкциялардан умумий йўқотишлар {fmt(c['jami_kw'], 1)} кВт")
    share = lambda x: f"{x['nomi']} – {fmt(x['kw'], 2)} кВт ({fmt(x['ulush'], 1)} %)"
    for i, key in zip(E["ulush"], ("tom", "devor", "deraza", "eshik")):
        set_text(els[i], share(by[key]))
    if "pol" in by:
        insert_after(els[E["ulush"][-1]], els[E["ulush"][-1]], [share(by["pol"])])

    # ── III. Xulosa
    rw0, rr0 = c["devor"]["R0"], c["tom"]["R0"]
    nw, nr = _num(d["norma_devor"]), _num(d["norma_tom"])
    all_ok = cmp_["devor"] and cmp_["tom"] and cmp_["deraza"] and cmp_["eshik"] and cmp_["pol"] is not False
    s = ["Бинонинг ташқи иссиқлик ҳимояси асосан меъёрий талабларга мос келади." if all_ok
         else "Бинонинг ташқи иссиқлик ҳимояси етарли даражада эмас."]
    wall_part = (f"ташқи деворларнинг иссиқлик узатилишига келтирилган қаршилиги {fmt(rw0, 2)} {U} бўлиб, "
                 + (f"меъёрий қийматдан ({nfmt(nw)} {U}) {_pct_below(rw0, nw)}%га паст." if not cmp_["devor"]
                    else f"меъёрий қийматга ({nfmt(nw)} {U}) мос келади."))
    s.append(f"Хусусан, фасадда иссиқлик ҳимоя қопламаси {d['fasad_izolyatsiya']}, {wall_part}")
    low = [n for n, ok in (("деразалар", cmp_["deraza"]), ("эшиклар", cmp_["eshik"])) if not ok]
    roof_part = (f"чордоқ ораёпмасининг қаршилиги {fmt(rr0, 2)} {U} бўлиб, "
                 + (f"меъёрий {nfmt(nr)} {U} га нисбатан {_pct_below(rr0, nr)}%га кам." if not cmp_["tom"]
                    else f"меъёрий {nfmt(nr)} {U} га мос келади."))
    if low:
        s.append(f"Шунингдек, {' ва '.join(low)}нинг қаршилиги ҳам меъёрий кўрсаткичлардан паст, {roof_part}")
    else:
        s.append(f"Деразалар ва эшикларнинг қаршилиги меъёрий кўрсаткичларга мос, {roof_part}")
    set_text(P("xulosa1"), " ".join(s))
    parts = [f"деворларидан {fmt(by['devor']['ulush'], 1)} %", f"томдан – {fmt(by['tom']['ulush'], 1)} %",
             f"деразалардан – {fmt(by['deraza']['ulush'], 1)} %"]
    if "pol" in by:
        parts.append(f"полдан – {fmt(by['pol']['ulush'], 1)} %")
    set_text(P("xulosa2"), f"Шу сабабли, бинодан иссиқлик йўқотилиши {', '.join(parts)}ни ташкил этади.")

    # ── IV. Tavsiyalar (4.1–4.4)
    if cmp_["devor"]:
        set_text(P("t_devor"), f"Ташқи деворларнинг амалдаги иссиқлик узатилишига қаршилиги {fmt(rw0, 2)} {U} бўлиб, меъёрий "
                               f"қиймат ({nfmt(nw)} {U}) га мос келади. Мавжуд қопламанинг ҳолатини сақлаш ва даврий кузатиб бориш тавсия этилади.")
    else:
        set_text(P("t_devor"), f"Ташқи деворларни ташқи томондан иссиқлик изоляция материаллари билан қоплаш тавсия этилади. "
                               f"Амалдаги қаршилик {fmt(rw0, 2)} {U} бўлиб, меъёрий қиймат {nfmt(nw)} {U} эканлигини ҳисобга "
                               "олган ҳолда, изоляция қалинлиги иссиқлик-техник ҳисоб асосида танланиши лозим.")
    if cmp_["tom"]:
        set_text(P("t_tom"), f"Чордоқ ораёпмасининг иссиқлик қаршилиги {fmt(rr0, 2)} {U}, меъёрий қиймати {nfmt(nr)} {U}. "
                             "Том конструкцияси меъёрий талабга мос; мавжуд иссиқлик изоляцияси қатламининг ҳолатини даврий текшириб бориш тавсия этилади.")
    else:
        set_text(P("t_tom"), f"Чордоқ ораёпмасининг иссиқлик қаршилиги {fmt(rr0, 2)} {U}, меъёрий қиймати эса {nfmt(nr)} {U}. "
                             "Шунинг учун том/чордоқ қисми орқали иссиқлик йўқотишларини камайтириш мақсадида мавжуд иссиқлик "
                             "изоляцияси қатламини текшириш ва зарур ҳолларда қўшимча минерал пахта, базальт плита ёки бошқа "
                             "самарали иссиқлик изоляцияси билан кучайтириш тавсия этилади.")
    rwin, rdoor = _num(d["deraza_R"]), _num(d["eshik_R"])
    ndw, ndd = _num(d["norma_deraza"]), _num(d["norma_eshik"])
    if cmp_["deraza"]:
        set_text(P("t_deraza"), f"Деразаларнинг иссиқлик қаршилиги {fmt(rwin, 2)} {U}, меъёрий қиймати {nfmt(ndw)} {U}. "
                                "Деразалар меъёрий талабга мос; зичловчи қистирмаларни даврий алмаштириб бориш тавсия этилади.")
    else:
        set_text(P("t_deraza"), f"Деразаларнинг иссиқлик қаршилиги {fmt(rwin, 2)} {U}, меъёрий қиймати {nfmt(ndw)} {U}. "
                                "Шунинг учун эски деразаларни икки ёки уч камерали энергия тежамкор ойнавандликка эга "
                                "конструкцияларга алмаштириш тавсия этилади.")
    if cmp_["eshik"]:
        set_text(P("t_eshik"), f"Кириш эшикларининг иссиқлик қаршилиги {fmt(rdoor, 2)} {U} бўлиб, меъёрий {nfmt(ndd)} {U} га мос. "
                               "Эшикларнинг зич ёпилишини таъминлаш (ёпгич, зичловчи) тавсия этилади.")
    else:
        set_text(P("t_eshik"), f"Кириш эшикларининг иссиқлик қаршилиги {fmt(rdoor, 2)} {U} бўлиб, меъёрий {nfmt(ndd)} {U} дан паст. "
                               "Эшикларни энергия самарадорлиги юқори бўлган, зич ёпиладиган конструкцияларга алмаштириш ёки "
                               "мавжуд эшикларни қўшимча иссиқлик изоляциялаш тавсия этилади.")

    # ── Rasmlar (surat yo'q bo'lsa, sarlavhasi bilan birga olib tashlanadi)
    for slot, pic, cap in (("rasm_1", "rasm1_pic", "rasm1_cap"), ("rasm_2", "rasm2_pic", "rasm2_cap"),
                           ("rasm_3", "rasm3_pic", "rasm3_cap")):
        if not (photos.get(slot) and _put_picture(document, P(pic), photos[slot], 15.5, 11.0)):
            remove(P(pic), P(cap))

    # Qolgan sariq belgilarni tozalash
    for hl in list(document.element.body.iter(qn("w:highlight"))):
        hl.getparent().remove(hl)

    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()
